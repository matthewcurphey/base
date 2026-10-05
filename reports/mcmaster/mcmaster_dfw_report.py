"""
McMaster DFW report: Dallas (DFW) McMaster batch analysis.

Not part of the daily pipeline — run manually (python run.py mcmaster-dfw-report).
Source list: reports/mcmaster/DFW_ORDERS.csv, an OBIEE export of every McMaster order line whose physical ship-to
is the Fort Worth, TX (DFW) site. Ship-to location isn't tracked anywhere in our
own Oracle/Castle feeds, so this file is the only way to identify these lines;
everything else (dates, job status, material tally) comes from our own data,
joined on org + so_nbr + so_line (+ shipment_nbr for the Castle join).

Verified before building this (against the 9/17/26 export):
- 1,466 rows in the file; 4 aren't real order lines (1 Quote that never became
  an order, 3 Credit Only adjustment lines) and are dropped by the join itself.
- The remaining 1,462 matched 100% against int_foundation_castle__sales_salesorder.
- "Shipped" per the file's Actual Ship Date agrees with our own invoice_date on
  1,461/1,462 — the one exception (SO 7387153) is a physically-shipped-but-not-
  yet-invoiced timing gap, not a data problem. We follow our own convention
  (invoice_date populated = shipped) for consistency with the rest of this report.
- All 348 currently-open lines were found in int_oracle__mcmaster_01_sales_production
  with real job_status data (100% match) — so the status-bucket logic below has
  what it needs.

Treats the Dallas batch as a fixed, gradually-loaded pool (order dates span
2026-04-29 to 2026-09-16): burndown (cumulative shipped vs remaining), not an
ongoing new/backlog/shipped trend like the regular McMaster contract.

Status buckets replicate int_oracle__mcmaster_02_open_backlog.sql's tally/floor/
status logic exactly (same partitioning, same window), with two changes:
  1. The open Dallas so_nbr/so_line combos are added back despite the 12/18
     placeholder promise date that excludes them from the live report.
  2. Within the tally's ORDER BY, Dallas lines are forced to sort after every
     other item of demand, regardless of what promise date they carry — "run
     the normal logic, then see if Dallas gets what's left over" rather than
     letting them compete on equal footing.
Everything else (the other ~1,100 non-Dallas 18-Dec-placeholder rows, all
normal in-scope backlog) is left exactly as the live report already computes it.
7 lines with a shipment split (so_shipment > 1) are excluded from status, same
as the live report already does for that class of line (comment there: "not
valid, causes duplicate fulfilment — surfaced in exception_to_cancel instead").
"""

import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import openpyxl
import pandas as pd

from etl.utils.connect_postgres import get_postgres_connection
from reports.mcmaster.mcmaster_output import (
    STATUS_ORDER,
    STATUS_COLORS,
    COLOR_BACKLOG,
    COLOR_NEW,
    COLOR_SHIPPED,
    COLOR_AXIS,
    COLOR_GRID,
    COLOR_TEXT_SECONDARY,
    range_size_px,
    _place_image,
    _aspect_figsize,
    build_backlog_status_pivot,
    _write_status_pivot,
)

DFW_CSV_PATH = os.path.join("reports", "mcmaster", "DFW_ORDERS.csv")
OUTPUT_PATH = os.path.join("reports", "mcmaster", "mcmaster_dfw_report.xlsx")
TREND_CHART_PATH = os.path.join("reports", "mcmaster", "mcmaster_dfw_burndown_trend.png")
STATUS_CHART_PATH = os.path.join("reports", "mcmaster", "mcmaster_dfw_status_chart.png")

# min_col, max_col, min_row, max_row — approved layout from the emailed
# version (resized in Excel from the first draft, which was skewed).
TREND_RANGE = (1, 6, 2, 26)
STATUS_CHART_RANGE = (12, 16, 2, 26)
COUNTS_PIVOT_TITLE_ROW = 30
COUNTS_PIVOT_HEADER_ROW = 31
COUNTS_PIVOT_COL = 1


def load_dallas_keys():
    """
    Real order lines from the DFW export:
    - drops Quotes / Credit Only lines, which aren't order/backlog lines at
      all (verified: every remaining row matches our own sales data)
    - drops Shipment Nbr > 1 (11 of 1,466 rows) — same rule the live report
      already applies everywhere else ("McMaster shipment splits are not
      valid and cause duplicate fulfilment"), so it's applied consistently
      here too rather than only at the status-bucket step. Of the 11: 7 are
      currently-open lines with no valid shipment-1 counterpart at all, 3 are
      already-shipped duplicates that would otherwise inflate the shipped
      count, and 1 (SO 7387153) has a separate, valid shipment-1 sibling that
      stays — only its shipment-2 duplicate is dropped.
    """
    df = pd.read_csv(DFW_CSV_PATH)
    df = df[df["Sales Type"].isin(["Order", "Invoice"])].copy()
    df = df[df["Shipment Nbr"].astype(str) == "1"].copy()
    for col in ("Sales Order Nbr", "Sales Line Nbr", "Shipment Nbr"):
        df[col] = df[col].astype(str)
    return df[["Inv Org Code", "Sales Order Nbr", "Sales Line Nbr", "Shipment Nbr"]].drop_duplicates()


def match_against_castle_sales(dallas_keys, engine):
    sales = pd.read_sql(
        """
        select inv_org_code, so_nbr, so_line, shipment_nbr, order_date, invoice_date
        from analytics_intermediate.int_foundation_castle__sales_salesorder
        """,
        engine,
    )
    for col in ("so_nbr", "so_line", "shipment_nbr"):
        sales[col] = sales[col].astype(str)

    matched = dallas_keys.merge(
        sales,
        left_on=["Inv Org Code", "Sales Order Nbr", "Sales Line Nbr", "Shipment Nbr"],
        right_on=["inv_org_code", "so_nbr", "so_line", "shipment_nbr"],
        how="inner",
    )
    return matched


def build_burndown(matched):
    """
    Daily cumulative-received/shipped/remaining-open burndown for the fixed
    batch, by org and total, from each line's real order_date/invoice_date —
    no dependency on historical ship-to tracking, since these dates are
    already stored per line regardless of whether the line currently passes
    the 18-Dec filter in the live report. Also carries same-day deltas
    (daily_received/daily_shipped) for the trend chart's flow areas.
    """
    matched = matched.copy()
    matched["order_date"] = pd.to_datetime(matched["order_date"])
    matched["invoice_date"] = pd.to_datetime(matched["invoice_date"])

    start = matched["order_date"].min().normalize()
    end = pd.Timestamp.today().normalize() - pd.Timedelta(days=1)
    spine = pd.date_range(start, end, freq="D")

    orgs = sorted(matched["inv_org_code"].unique())
    frames = []
    for org in orgs + ["Total"]:
        sub = matched if org == "Total" else matched[matched["inv_org_code"] == org]
        total_n = len(sub)
        received = [(sub["order_date"] <= dt).sum() for dt in spine]
        shipped = [(sub["invoice_date"].notna() & (sub["invoice_date"] <= dt)).sum() for dt in spine]
        org_df = pd.DataFrame({"dt": spine, "org": org, "total_lines": total_n,
                                "cum_received": received, "cum_shipped": shipped})
        org_df["remaining_open"] = org_df["cum_received"] - org_df["cum_shipped"]
        org_df["daily_received"] = org_df["cum_received"].diff().fillna(org_df["cum_received"].iloc[0])
        org_df["daily_shipped"] = org_df["cum_shipped"].diff().fillna(org_df["cum_shipped"].iloc[0])
        frames.append(org_df)
    return pd.concat(frames, ignore_index=True)


def build_status_buckets(matched, engine):
    """
    Replicates int_oracle__mcmaster_02_open_backlog.sql's tally/floor/status
    CTEs directly in SQL (same partitioning/window logic), but:
      - re-includes the verified open Dallas so_nbr/so_line combos despite the
        12/18 placeholder that excludes them from the live report
      - forces those same lines to sort last within each item's tally via a
        synthetic max sort key, so they only ever pick up leftover material
    Every other row (normal in-scope backlog, and the other 18-Dec lines that
    are NOT part of this Dallas batch) is filtered/ordered exactly as the live
    report already does — nothing else changes.
    """
    open_lines = matched[matched["invoice_date"].isna()][["inv_org_code", "so_nbr", "so_line"]].drop_duplicates()
    pairs = list(open_lines[["so_nbr", "so_line"]].itertuples(index=False, name=None))
    print(f"  {len(pairs)} distinct open Dallas (so_nbr, so_line) combos to re-include")

    # A literal VALUES list is simplest and safe here — so_nbr/so_line are
    # digit-string identifiers from our own verified join, not user input.
    values_sql = ",".join(f"('{so}','{line}')" for so, line in pairs)

    query = f"""
    with dallas_keys(so_nbr, so_line) as (
        values {values_sql}
    ),

    base as (
        select
            b.*,
            (b.so_nbr, b.so_line) in (select so_nbr, so_line from dallas_keys) as is_dallas_open
        from analytics_intermediate.int_oracle__mcmaster_01_sales_production b
        where (b.is_mcmaster or b.comp_inv_req > 0)
          and not (b.is_mcmaster and b.so_shipment::int > 1)
          and (
              b.promise_date is distinct from '2026-12-18'
              or (b.so_nbr, b.so_line) in (select so_nbr, so_line from dallas_keys)
          )
    ),

    tallied as (
        select
            *,
            inv_atl - sum(case when inv_org_code = 'ATL' then coalesce(comp_inv_req, 0) else 0 end)
                over (partition by item_clean order by is_dallas_open, promise_date nulls last, so_nbr, so_line
                      rows unbounded preceding) as tally_atl,
            inv_cle - sum(case when inv_org_code = 'CLE' then coalesce(comp_inv_req, 0) else 0 end)
                over (partition by item_clean order by is_dallas_open, promise_date nulls last, so_nbr, so_line
                      rows unbounded preceding) as tally_cle,
            inv_dal - sum(case when inv_org_code = 'DAL' then coalesce(comp_inv_req, 0) else 0 end)
                over (partition by item_clean order by is_dallas_open, promise_date nulls last, so_nbr, so_line
                      rows unbounded preceding) as tally_dal,
            inv_jvl - sum(case when inv_org_code = 'JVL' then coalesce(comp_inv_req, 0) else 0 end)
                over (partition by item_clean order by is_dallas_open, promise_date nulls last, so_nbr, so_line
                      rows unbounded preceding) as tally_jvl,
            inv_los - sum(case when inv_org_code = 'LOS' then coalesce(comp_inv_req, 0) else 0 end)
                over (partition by item_clean order by is_dallas_open, promise_date nulls last, so_nbr, so_line
                      rows unbounded preceding) as tally_los,
            inv_wie - sum(case when inv_org_code = 'WIE' then coalesce(comp_inv_req, 0) else 0 end)
                over (partition by item_clean order by is_dallas_open, promise_date nulls last, so_nbr, so_line
                      rows unbounded preceding) as tally_wie
        from base
    ),

    floored as (
        select
            *,
            case inv_org_code
                when 'ATL' then tally_atl when 'CLE' then tally_cle when 'DAL' then tally_dal
                when 'JVL' then tally_jvl when 'LOS' then tally_los when 'WIE' then tally_wie
            end as tally_home_org,
            case inv_org_code
                when 'ATL' then tally_atl < 0 and comp_inv_req > 0
                when 'CLE' then tally_cle < 0 and comp_inv_req > 0
                when 'DAL' then tally_dal < 0 and comp_inv_req > 0
                when 'JVL' then tally_jvl < 0 and comp_inv_req > 0
                when 'LOS' then tally_los < 0 and comp_inv_req > 0
                when 'WIE' then tally_wie < 0 and comp_inv_req > 0
            end as is_short
        from tallied
    ),

    statused as (
        select
            *,
            case
                when not is_mcmaster then null
                when comp_inv_req = 0
                     and job_status ~ '^(Complete|Closed)(,(Complete|Closed))*$' then 'Job Complete'
                when comp_inv_req = 0 then 'Job Started'
                when is_short then 'No Material'
                when comp_inv_req > 0 then 'Material Available'
                else 'Unknown'
            end as mcm_status
        from floored
    )

    select inv_org_code, so_nbr, so_line, item_clean, dj_nbr, job_status, mcm_status,
           total_sales_usd, is_dallas_open
    from statused
    where is_dallas_open
    """
    return pd.read_sql(query, engine)


def build_burndown_chart(burndown_df, output_path, figsize=(9.6, 8.1), dpi=200):
    """Company-wide burndown — same visual language as the live report's trend
    chart (grey level bar, red/green flow areas straddling zero), but the bar
    is remaining-open (a level) and the areas are daily received/shipped
    (flows) across the whole batch history, not a rolling N-day window."""
    df = burndown_df[burndown_df["org"] == "Total"].sort_values("dt")
    start_date = df["dt"].min()

    fig, ax = plt.subplots(figsize=figsize, dpi=dpi)
    fig.patch.set_facecolor("#fcfcfb")
    ax.set_facecolor("#fcfcfb")

    ax.bar(df["dt"], df["remaining_open"], color=COLOR_BACKLOG, width=0.9, label="Remaining Open", zorder=2)
    ax.fill_between(df["dt"], 0, df["daily_received"], color=COLOR_NEW, alpha=0.85, linewidth=0,
                     label="Received", zorder=3)
    ax.fill_between(df["dt"], 0, -df["daily_shipped"], color=COLOR_SHIPPED, alpha=0.85, linewidth=0,
                     label="Shipped", zorder=3)

    ax.axhline(0, color=COLOR_AXIS, linewidth=1)
    ax.grid(axis="y", color=COLOR_GRID, linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color(COLOR_AXIS)

    ax.tick_params(colors=COLOR_TEXT_SECONDARY, labelsize=9)
    ax.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=mdates.MO, interval=2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d-%b"))
    fig.autofmt_xdate(rotation=0, ha="center")

    fig.suptitle(
        f"DFW Batch - Backlog Burndown Since {start_date:%d-%b-%Y}",
        fontsize=13, color="#0b0b0b", x=0.01, ha="left", y=0.98,
    )
    handles, labels = ax.get_legend_handles_labels()
    order = [labels.index(name) for name in ("Remaining Open", "Received", "Shipped")]
    ax.legend(
        [handles[i] for i in order], [labels[i] for i in order],
        loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=3,
        frameon=False, fontsize=9, labelcolor=COLOR_TEXT_SECONDARY,
    )

    fig.tight_layout(rect=(0, 0.04, 1, 0.93))
    fig.savefig(output_path, facecolor=fig.get_facecolor())
    plt.close(fig)
    return output_path


def build_status_waterfall_chart(status_df, output_path, figsize=(9.6, 4.8), dpi=200):
    """Company-wide waterfall — same design as the live report's status chart:
    a grey Total bar, then each stage below it offset by the running total."""
    totals = status_df["mcm_status"].value_counts()
    totals = totals.reindex([s for s in STATUS_ORDER if s in totals.index]).fillna(0).astype(int)
    grand_total = totals.sum()
    cum_starts = totals.cumsum().shift(fill_value=0)

    labels = ["Total Open"] + list(totals.index)
    values = [grand_total] + list(totals.values)
    lefts = [0] + list(cum_starts.values)
    colors = [COLOR_BACKLOG] + [STATUS_COLORS[s] for s in totals.index]

    fig, ax = plt.subplots(figsize=figsize, dpi=dpi)
    fig.patch.set_facecolor("#fcfcfb")
    ax.set_facecolor("#fcfcfb")

    GAP = 0.9
    y_pos = [0] + [1 + GAP + i for i in range(len(totals))]
    bar_heights = [0.7] + [0.6] * len(totals)
    bars = ax.barh(y_pos, values, left=lefts, color=colors, height=bar_heights, zorder=2)
    ax.set_yticks(y_pos)
    ax.set_yticklabels(labels, fontsize=10, color="#0b0b0b")
    ax.get_yticklabels()[0].set_fontweight("bold")
    ax.invert_yaxis()

    for bar, value in zip(bars, values):
        ax.text(
            bar.get_x() + bar.get_width() + grand_total * 0.02, bar.get_y() + bar.get_height() / 2,
            f"{int(value):,}", va="center", ha="left", fontsize=10, color="#0b0b0b",
        )

    ax.set_xlim(0, grand_total * 1.15)
    ax.set_xticks([])
    for spine in ("top", "right", "bottom", "left"):
        ax.spines[spine].set_visible(False)

    fig.suptitle(
        "DFW Batch - Open Backlog by Status (Company-Wide)", fontsize=13, color="#0b0b0b",
        x=0.01, ha="left", y=0.98,
    )

    fig.tight_layout(rect=(0, 0.02, 1, 0.90))
    fig.savefig(output_path, facecolor=fig.get_facecolor())
    plt.close(fig)
    return output_path


def _style_pivot_block(ws, pivot, header_row, col, number_format):
    """Light, self-contained styling for a pivot written into a brand-new
    workbook (no existing hand-formatted template to inherit from here) —
    bold header row + Total row/column, thin borders, sensible number format."""
    from openpyxl.styles import Border, Font, PatternFill, Side

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="EFEFEA")
    total_fill = PatternFill("solid", fgColor="F5F5F0")

    n_cols = len(pivot.columns)
    n_rows = len(pivot.index)
    for row_offset in range(n_rows + 1):
        row = header_row + row_offset
        is_header = row_offset == 0
        is_total_row = row_offset == n_rows  # margins puts "Total" last
        for col_offset in range(n_cols + 1):
            cell = ws.cell(row, col + col_offset)
            cell.border = border
            if is_header or col_offset == 0:
                cell.font = Font(bold=True)
            if is_header:
                cell.fill = header_fill
            elif is_total_row:
                cell.fill = total_fill
                cell.font = Font(bold=True)
            if not is_header and col_offset > 0:
                cell.number_format = number_format


def main():
    print("Loading DFW orders export and matching against our sales data...")
    dallas_keys = load_dallas_keys()

    engine = get_postgres_connection()
    matched = match_against_castle_sales(dallas_keys, engine)
    print(f"  {len(matched)} lines matched (of {len(dallas_keys)} real order lines in the file)")

    print("Building historical burndown...")
    burndown = build_burndown(matched)

    print("Building current-status buckets for the open Dallas lines...")
    status = build_status_buckets(matched, engine)
    engine.dispose()

    print("\nStatus bucket summary (open Dallas lines):")
    print(status["mcm_status"].value_counts())

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)

    print("Writing workbook...")
    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        # A blank placeholder sheet — charts/pivots get added to it below via
        # openpyxl directly, after this writer closes and re-saves the file.
        pd.DataFrame().to_excel(writer, sheet_name="summary", index=False)
        burndown.to_excel(writer, sheet_name="burndown_trend", index=False)
        status.to_excel(writer, sheet_name="open_status", index=False)

    wb = openpyxl.load_workbook(OUTPUT_PATH)
    ws = wb["summary"]

    last_dt = burndown["dt"].max()
    ws["A1"] = f"Historical View - Burndown Through {last_dt:%d-%b-%Y}"
    ws["M1"] = "Current View - Open Lines by Status"

    trend_w, trend_h = range_size_px(ws, *TREND_RANGE)
    trend_final = build_burndown_chart(burndown, TREND_CHART_PATH, figsize=_aspect_figsize(trend_w, trend_h))
    _place_image(ws, trend_final, "A2", *TREND_RANGE)

    status_w, status_h = range_size_px(ws, *STATUS_CHART_RANGE)
    status_final = build_status_waterfall_chart(status, STATUS_CHART_PATH, figsize=_aspect_figsize(status_w, status_h))
    _place_image(ws, status_final, "L2", *STATUS_CHART_RANGE)

    counts_df = status.assign(line_count=1)
    counts_pivot = build_backlog_status_pivot(counts_df, value_col="line_count")

    ws.cell(COUNTS_PIVOT_TITLE_ROW, COUNTS_PIVOT_COL, "Open Lines by Org & Status")
    ws.cell(COUNTS_PIVOT_TITLE_ROW, COUNTS_PIVOT_COL).font = openpyxl.styles.Font(bold=True)
    _write_status_pivot(ws, counts_pivot, COUNTS_PIVOT_HEADER_ROW, COUNTS_PIVOT_COL, cast=int)
    _style_pivot_block(ws, counts_pivot, COUNTS_PIVOT_HEADER_ROW, COUNTS_PIVOT_COL, "#,##0")

    for col_letter, width in (("A", 10), ("B", 14), ("C", 12), ("D", 12), ("E", 15), ("F", 10)):
        ws.column_dimensions[col_letter].width = width

    wb.save(OUTPUT_PATH)
    print(f"\nWritten: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
